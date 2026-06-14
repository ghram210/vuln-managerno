from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_cv_template():
    doc = Document()

    # Set page margins
    section = doc.sections[0]
    section.left_margin = Inches(3.0)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.header_distance = Inches(0)

    # Access the header
    header = section.header
    header_para = header.paragraphs[0]

    # XML for background shape
    vml_xml = r"""
    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">
        <w:pict>
            <v:rect id="sidebar_bg" style="position:absolute;margin-left:-216pt;margin-top:-36pt;width:180pt;height:842pt;z-index:-1" fillcolor="#70924e" stroked="f">
                <v:fill opacity="1"/>
            </v:rect>
        </w:pict>
    </w:r>
    """
    header_para._p.append(parse_xml(vml_xml))

    # XML for sidebar text
    sidebar_text_xml = r"""
    <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">
        <w:pict>
            <v:shape id="sidebar_text" style="position:absolute;margin-left:-200pt;margin-top:20pt;width:150pt;height:700pt;z-index:1">
                <v:textbox>
                    <w:txbxContent>
                        <w:p>
                            <w:r><w:rPr><w:b/><w:sz w:val="32"/><w:color w:val="FFFFFF"/></w:rPr><w:t>CONTACT</w:t></w:r>
                        </w:p>
                        <w:p>
                            <w:r><w:rPr><w:sz w:val="20"/><w:color w:val="FFFFFF"/></w:rPr><w:t>Address: City, Country</w:t></w:r>
                        </w:p>
                        <w:p/>
                        <w:p>
                            <w:r><w:rPr><w:b/><w:sz w:val="32"/><w:color w:val="FFFFFF"/></w:rPr><w:t>ACADEMIC</w:t></w:r>
                        </w:p>
                        <w:p>
                            <w:r><w:rPr><w:sz w:val="20"/><w:color w:val="FFFFFF"/></w:rPr><w:t>Degree Name</w:t></w:r>
                        </w:p>
                    </w:txbxContent>
                </v:textbox>
            </v:shape>
        </w:pict>
    </w:r>
    """
    header_para._p.append(parse_xml(sidebar_text_xml))

    # Main content
    title = doc.add_paragraph()
    run = title.add_run("Your Name Here")
    run.font.size = Pt(28)
    run.font.bold = True

    doc.add_paragraph()

    p_heading = doc.add_paragraph()
    run = p_heading.add_run("PROFILE")
    run.font.size = Pt(14)
    run.font.bold = True

    doc.add_paragraph("This is your profile summary. Since the sidebar is in the header, it will automatically appear on every new page. Press Enter to reach a new page.")

    doc.add_paragraph()

    c_heading = doc.add_paragraph()
    run = c_heading.add_run("CAREER")
    run.font.size = Pt(14)
    run.font.bold = True

    doc.add_paragraph("Job Title | 2020 - Present")
    doc.add_paragraph("• Responsibility 1\n• Responsibility 2")

    doc.add_page_break()
    doc.add_paragraph("Second page - sidebar should be here too.")

    doc.save("CV_Template.docx")
    print("CV_Template.docx created.")

if __name__ == "__main__":
    create_cv_template()
