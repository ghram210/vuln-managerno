from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

def set_rtl(paragraph):
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = pPr.makeelement(qn('w:bidi'))
        pPr.append(bidi)
    bidi.set(qn('w:val'), '1')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def add_arabic_paragraph(doc, text, is_heading=False, level=1):
    if is_heading:
        p = doc.add_heading(text, level=level)
    else:
        p = doc.add_paragraph(text)
    set_rtl(p)
    return p

doc = Document()

# Global Style
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(12)

# Title
add_arabic_paragraph(doc, 'الفصل الرابع: الجانب التطبيقي والهندسة البرمجية', is_heading=True, level=0)

# Section 4.1 Outline
add_arabic_paragraph(doc, '4.1 مقدمة الفصل العملي', is_heading=True, level=1)
add_arabic_paragraph(doc, "يتناول هذا الفصل الترجمة الفعلية للمفاهيم النظرية إلى واقع برمجي، حيث يتم استعراض هندسة المنصة وكيفية عملها في بيئة Kali Linux للوصول إلى نظام مؤتمت بالكامل لاكتشاف وإدارة الثغرات.")

add_arabic_paragraph(doc, '4.2 مخطط هيكلية الفصل العملي (Proposed Outline)', is_heading=True, level=1)
outline_points = [
    "4.2.1 إعداد بيئة المختبر (Kali Linux Environment Setup).",
    "4.2.2 هندسة المحرك التنفيذي (Backend Orchestration): الربط بين FastAPI والأدوات الأمنية.",
    "4.2.3 منطق أتمتة الاستخبارات (Intelligence Pipeline): كيف تتحول النتائج الخام إلى معلومات أمنية.",
    "4.2.4 معالجة البيانات وتدفقها (Data Flow Architecture): من سطر الأوامر إلى قاعدة البيانات.",
    "4.2.5 دليل واجهات الاستخدام (UI Operational Guide): استعراض عملي للمنصة مع الصور.",
    "4.2.6 تحليل النتائج وإصدار التقارير (Reporting & Analytics)."
]
for point in outline_points:
    add_arabic_paragraph(doc, point)

# Section 4.3 UI Flow for Screenshots
add_arabic_paragraph(doc, '4.3 دليل تدفق واجهات المنصة (لإضافة الصور)', is_heading=True, level=1)
add_arabic_paragraph(doc, "فيما يلي النص المخصص لوصف عمل الواجهات بالترتيب الصحيح ليتم وضعه بجانب صور الشاشة الفعلية:")

ui_flow = [
    "1. واجهة الدخول: تبدأ الدورة التشغيلية عبر بوابة وصول آمنة تضمن عزل بيانات كل مختبر أمني.",
    "2. لوحة التحكم: تعرض ملخصاً لحظياً لنشاط الأصول وتوزيع المخاطر المكتشفة بناءً على خطورتها.",
    "3. إدارة النطاقات: يتم هنا إضافة المواقع المستهدفة والتحقق من ملكيتها لضمان شرعية الفحص.",
    "4. واجهة إطلاق الفحص: تتيح تخصيص المهمة عبر اختيار الأدوات الأمنية المناسبة وتحديد عمق المسح.",
    "5. مراقبة التنفيذ: واجهة تفاعلية تعرض حالة الأدوات وتدفق مخرجات سطر الأوامر (Terminal) في الوقت الفعلي.",
    "6. سجل الثغرات المكتشفة: قاعدة بيانات مهيكلة تعرض الثغرات بعد ربطها آلياً بمعرفات CVE وأكواد الاستغلال.",
    "7. تحليل ناقل الهجوم: تقدم رؤية بصرية عميقة لمكونات كل ثغرة وتأثيرها على سرية وسلامة البيانات.",
    "8. التقارير الإدارية: المرحلة النهائية حيث يتم تصدير ملخص تنفيذي احترافي لصناع القرار."
]

for step in ui_flow:
    add_arabic_paragraph(doc, step)

doc.save('Chapter_4_Practical_Implementation.docx')
